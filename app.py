import os, io, json, base64, re
from datetime import date
import streamlit as st
from PIL import Image
from openai import OpenAI
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

st.set_page_config(page_title="Garment QA CAPA", page_icon="🧵", layout="wide")
BLUE = RGBColor(0,112,192); DARK = RGBColor(30,30,30)

def data_url(f):
    return f"data:{f.type};base64," + base64.b64encode(f.getvalue()).decode()

def analyze(f, api_key):
    client = OpenAI(api_key=api_key)
    prompt = """You are a professional garment factory QA inspector. Analyze this defect photo.
Return ONLY valid JSON with keys: defect, root_cause, corrective_action, severity.
Use concise professional QA wording. Describe only what is visibly supported.
For root cause, give the most likely process/machine/material cause; if uncertain say
'To be confirmed by factory investigation.' Corrective action must be practical and include
recheck/prevention. Severity must be Minor, Major, or Critical. Do not invent specifications."""
    r = client.responses.create(model="gpt-5.6-luna", input=[{"role":"user","content":[
        {"type":"input_text","text":prompt},{"type":"input_image","image_url":data_url(f)}
    ]}])
    txt = re.sub(r"^```(?:json)?\s*|\s*```$","",r.output_text.strip()).strip()
    return json.loads(txt)

def txt(slide, text, x,y,w,h,size=12,bold=False):
    b=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
    p=b.text_frame.paragraphs[0]; p.alignment=PP_ALIGN.CENTER
    r=p.add_run(); r.text=text; r.font.name="Arial"; r.font.size=Pt(size); r.font.bold=bold; r.font.color.rgb=DARK

def bar(slide,text,x,y,w,h):
    s=slide.shapes.add_shape(1,Inches(x),Inches(y),Inches(w),Inches(h))
    s.fill.solid(); s.fill.fore_color.rgb=BLUE; s.line.color.rgb=BLUE
    txt(slide,text,x,y,w,h,12,True)

def add_image_fit(slide, data, x,y,w,h):
    im=Image.open(io.BytesIO(data))
    scale=min(w/im.width,h/im.height)
    ww=im.width*scale/96; hh=im.height*scale/96
    slide.shapes.add_picture(io.BytesIO(data), Inches(x+(w-ww)/2), Inches(y+(h-hh)/2), width=Inches(ww), height=Inches(hh))

def make_ppt(factory, report_date, defect_file, evidence_file, result):
    prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
    slide=prs.slides.add_slide(prs.slide_layouts[6])
    hdr=slide.shapes.add_shape(1,0,0,Inches(13.333),Inches(.68))
    hdr.fill.solid(); hdr.fill.fore_color.rgb=BLUE; hdr.line.color.rgb=BLUE
    txt(slide,f"FACTORY: {factory}",.1,.03,13.13,.27,18,True)
    txt(slide,f"DATE: {report_date}",.1,.32,13.13,.25,14,True)
    cols=[(0,4.44,"DEFECT IMAGES"),(4.44,4.44,"DEFECT"),(8.88,4.45,"EVIDENCE IMAGES")]
    for x,w,label in cols:
        bar(slide,label,x,.75,w,.36)
        r=slide.shapes.add_shape(1,Inches(x),Inches(1.11),Inches(w),Inches(6.29))
        r.fill.background(); r.line.color.rgb=RGBColor(120,120,120)
    add_image_fit(slide,defect_file.getvalue(),.15,1.25,4.14,5.9)
    txt(slide,result["defect"],4.62,1.45,4.08,1.15,14)
    bar(slide,"CAP",4.44,4.22,4.44,.36)
    cap=f"ROOT CAUSE: {result['root_cause']}\\n\\nCORRECTIVE ACTION: {result['corrective_action']}"
    txt(slide,cap,4.62,4.72,4.08,1.85,10)
    if evidence_file:
        add_image_fit(slide,evidence_file.getvalue(),9.03,1.25,4.15,5.9)
    else:
        txt(slide,"Evidence image: Not provided",9.1,3.65,3.95,.6,12)
    out=io.BytesIO(); prs.save(out); return out.getvalue()

st.title("🧵 Garment QA — Defect → CAPA → PPT")
st.caption("First-level prototype for garment inspection reporting.")

with st.sidebar:
    st.header("Report Details")
    factory=st.text_input("Factory","HOA VU (GUOTAI)")
    report_date=st.date_input("Date",date.today())
    api_key=st.text_input("OpenAI API key",type="password")
    st.caption("You can also set OPENAI_API_KEY in the environment.")

defect_file=st.file_uploader("1. Upload defective image",type=["jpg","jpeg","png","webp"])
evidence_file=st.file_uploader("2. Optional evidence / corrected image",type=["jpg","jpeg","png","webp"])

if defect_file:
    st.image(defect_file,caption="Defect image",width=420)
    if st.button("🔎 Analyze Defect",type="primary"):
        key=api_key or os.getenv("OPENAI_API_KEY")
        if not key: st.error("Please enter an OpenAI API key.")
        else:
            with st.spinner("Analyzing defect..."):
                try: st.session_state.result=analyze(defect_file,key)
                except Exception as e: st.error(f"Analysis failed: {e}")

if "result" in st.session_state:
    r=st.session_state.result
    st.subheader("AI Draft — review before release")
    r["defect"]=st.text_area("Defect",r.get("defect",""))
    r["root_cause"]=st.text_area("Root Cause",r.get("root_cause",""))
    r["corrective_action"]=st.text_area("Corrective Action",r.get("corrective_action",""))
    r["severity"]=st.selectbox("Severity",["Minor","Major","Critical"],index=["Minor","Major","Critical"].index(r.get("severity","Major")) if r.get("severity") in ["Minor","Major","Critical"] else 1)
    if st.button("📊 Generate PowerPoint",type="primary"):
        data=make_ppt(factory,report_date.strftime("%m/%d/%Y"),defect_file,evidence_file,r)
        name=f"QA_CAPA_{factory.replace(' ','_')}_{report_date.strftime('%Y%m%d')}.pptx"
        st.download_button("⬇️ Download PPT",data,name,"application/vnd.openxmlformats-officedocument.presentationml.presentation")
else:
    st.info("Upload a defect image to start.")
